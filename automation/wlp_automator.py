#!/usr/bin/env python3
"""
wlp_automator.py — drive the Academic WLP web app through its workload tabs.

Headed Chromium automation built with Playwright (sync API). It reads a YAML
config (config/tabs.yaml by default) describing which tab to fill in, which
local file(s) to upload, which button to click, and which DOM element to
watch for to know the result is ready.

Usage:
    python wlp_automator.py                  # run with default config
    python wlp_automator.py --config my.yaml # custom config
    python wlp_automator.py --dry-run        # print plan, do not launch browser

Designed to stay simple and explicit so you can read the code and know exactly
what it is going to click. Nothing here touches the WLP app's source — all
interactions go through the browser like a human would.
"""

from __future__ import annotations

import argparse
import logging
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from playwright.sync_api import (
    Locator,
    Page,
    TimeoutError as PWTimeoutError,
    sync_playwright,
)

LOG = logging.getLogger("wlp")


def _normalise_steps(raw: Any) -> list[dict[str, str]]:
    """Accept preflight/postflight as either:
       - a list of {"check": sel} / {"click": sel} dicts, or
       - a list of bare selectors (treated as {"check": sel}).
    Returns the canonical list-of-dicts form.
    """
    if not raw:
        return []
    out: list[dict[str, str]] = []
    for item in raw:
        if isinstance(item, str):
            out.append({"check": item})
        elif isinstance(item, dict):
            out.append({k: v for k, v in item.items() if k in ("check", "click")})
        else:
            raise ValueError(f"preflight/postflight entries must be str or dict, got {type(item)}")
    return [s for s in out if s]


def run_steps(page: Page, steps: list[dict[str, str]], phase: str) -> None:
    """Execute a sequence of {"check": sel} / {"click": sel} steps."""
    for step in steps:
        for action, selector in step.items():
            if action == "check":
                LOG.info("  → %s check %s", phase, selector)
                loc = page.locator(selector)
                if not loc.is_checked():
                    loc.check(force=True)
            elif action == "click":
                LOG.info("  → %s click %s", phase, selector)
                page.locator(selector).click(force=True, timeout=5000)
            else:
                LOG.warning("  unknown %s action '%s'", phase, action)


# ──────────────────────────────────────────────────────────────────────────
# Config model
# ──────────────────────────────────────────────────────────────────────────


@dataclass
class TabSpec:
    key: str
    enabled: bool
    main_tab: str
    sub_tab: str
    file_input: str
    action: str
    wait_for: str
    files: list[Path]
    notes: str = ""
    # Optional 'kind' field. Default behaviour is "upload" — one or more files
    # dropped into a single <input type="file">, then a button click and wait.
    # 'citizenship_paste' is special: each entry under `paste:` is an XLSX that
    # is converted to TSV and dropped into the corresponding SharePoint-paste
    # textarea (teaching/research/school); then a single Analyse click merges
    # all three with cross-category dedup.
    kind: str = "upload"
    paste: dict[str, Path] = field(default_factory=dict)
    # Optional preflight / postflight hooks. Each is a list of actions,
    # each action a {"check": selector} or {"click": selector} dict.
    # preflight runs after file upload but before the action button click.
    # postflight runs after the action click but before waiting for results.
    preflight: list[dict[str, str]] = field(default_factory=list)
    postflight: list[dict[str, str]] = field(default_factory=list)
    # Optional fields used by kind="combined_tag_setup": define a new tag rule
    # (rule_name with rule_reduction_pct % reduction) and apply it to the row
    # whose displayed name matches apply_to_name.
    rule_name: str = ""
    rule_reduction_pct: int = 0
    apply_to_name: str = ""
    # If true, tick the 'Teaching staff only' checkbox (#combTlOnly) before
    # searching for the named academic — useful when the target row appears
    # in teaching data and you want to narrow the visible table.
    teaching_staff_only: bool = False
    # Optional fields used by kind="combined_tag_assignments": read an XLSX
    # file with one row per academic, columns for each tag, and apply the
    # tags whose cells match `assignments_apply_marker` (default "Y"). The
    # tag rules themselves must already be defined (e.g. by a preceding
    # `combined_tag_setup` entry) — this kind only applies existing tags,
    # it does not define new ones. A tag without a matching rule is still
    # stored on the academic but has no workload effect.
    assignments_xlsx: str = ""
    assignments_sheet: str = ""
    assignments_name_column: int = 1
    assignments_apply_marker: str = "Y"
    assignments_tag_columns: dict[str, int] = field(default_factory=dict)

    @classmethod
    def from_dict(cls, d: dict[str, Any], data_dir: Path) -> "TabSpec":
        files_raw = d.get("files") or []
        if isinstance(files_raw, str):
            files_raw = [files_raw]
        files = [data_dir / f for f in files_raw]
        paste_raw = d.get("paste") or {}
        paste = {cat: data_dir / fname for cat, fname in paste_raw.items()}
        preflight = _normalise_steps(d.get("preflight"))
        postflight = _normalise_steps(d.get("postflight"))
        return cls(
            key=d["key"],
            enabled=bool(d.get("enabled", True)),
            main_tab=d["main_tab"],
            sub_tab=d.get("sub_tab", "") or "",
            file_input=d.get("file_input", "") or "",
            action=d.get("action", "") or "",
            wait_for=d.get("wait_for", "") or "",
            files=files,
            notes=d.get("notes", "") or "",
            kind=d.get("kind", "upload"),
            paste=paste,
            preflight=preflight,
            postflight=postflight,
            rule_name=d.get("rule_name", "") or "",
            rule_reduction_pct=int(d.get("rule_reduction_pct", 0) or 0),
            apply_to_name=d.get("apply_to_name", "") or "",
            teaching_staff_only=bool(d.get("teaching_staff_only", False)),
            assignments_xlsx=d.get("assignments_xlsx", "") or "",
            assignments_sheet=d.get("assignments_sheet", "") or "",
            assignments_name_column=int(d.get("assignments_name_column", 1) or 1),
            assignments_apply_marker=d.get("assignments_apply_marker", "Y") or "Y",
            assignments_tag_columns=dict(d.get("assignments_tag_columns", {}) or {}),
        )


@dataclass
class AppConfig:
    url: str
    headless: bool
    slow_mo_ms: int
    result_timeout_s: int
    screenshot_dir: str
    data_dir: Path
    tabs: list[TabSpec] = field(default_factory=list)

    @classmethod
    def load(cls, path: Path) -> "AppConfig":
        raw = yaml.safe_load(path.read_text())
        app = raw.get("app", {})
        data_dir = Path(raw["data_dir"]).expanduser().resolve()
        tabs = [TabSpec.from_dict(t, data_dir) for t in raw.get("tabs", [])]
        return cls(
            url=app.get("url", ""),
            headless=bool(app.get("headless", False)),
            slow_mo_ms=int(app.get("slow_mo_ms", 0)),
            result_timeout_s=int(app.get("result_timeout_s", 30)),
            screenshot_dir=app.get("screenshot_dir", "") or "",
            data_dir=data_dir,
            tabs=tabs,
        )


# ──────────────────────────────────────────────────────────────────────────
# Browser helpers
# ──────────────────────────────────────────────────────────────────────────


def banner_text(page: Page) -> str:
    """Returns the visible text of the session-restore banner, or '' if hidden."""
    return page.evaluate(
        "() => { const b = document.getElementById('sessionBanner');"
        " return (b && getComputedStyle(b).display !== 'none') ? b.innerText.trim() : ''; }"
    )


def dismiss_session_banner(page: Page) -> None:
    """If the app shows a 'restore previous session?' banner, click Clear."""
    text = banner_text(page)
    if not text:
        return
    LOG.info("Session banner detected — clearing previous session")
    page.locator("#sessionClear").click()
    # Give the app a moment to clear localStorage and re-render.
    page.wait_for_timeout(300)


def switch_main_tab(page: Page, main_group: str) -> None:
    """Click the top-level tab whose data-maingroup matches."""
    sel = f'.main-tab-btn[data-maingroup="{main_group}"]'
    page.locator(sel).click()
    # Sub-tab bars swap on the same animation frame; this just gives the
    # browser a tick to update DOM before we click the sub-tab.
    page.wait_for_timeout(150)


def switch_sub_tab(page: Page, panel_id: str) -> None:
    """Click the sub-tab whose data-panel matches (e.g. 'panel-project')."""
    sel = f'.sub-tab-btn[data-panel="{panel_id}"]'
    page.locator(sel).click()
    page.wait_for_timeout(150)


def wait_for_action_enabled(page: Page, selector: str, timeout_s: int) -> None:
    """Wait until the analyse button is enabled and visible."""
    LOG.info("Waiting for %s to be enabled (timeout %ds)", selector, timeout_s)
    deadline = time.monotonic() + timeout_s
    n_polls = 0
    while time.monotonic() < deadline:
        ok = page.evaluate(
            "sel => { const el = document.querySelector(sel);"
            " return !!(el && !el.disabled && getComputedStyle(el).display !== 'none'); }",
            arg=selector,
        )
        n_polls += 1
        if ok:
            LOG.info("  enabled after %d polls", n_polls)
            return
        page.wait_for_timeout(150)
    state = page.evaluate(
        "sel => { const el=document.querySelector(sel);"
        " return el ? `disabled=${el.disabled} display=${getComputedStyle(el).display}` : 'missing'; }",
        arg=selector,
    )
    LOG.info("  timed out after %d polls; final state: %s", n_polls, state)
    raise PWTimeoutError(f"{selector} did not become enabled within {timeout_s}s")


def wait_for_result(page: Page, selector: str, timeout_s: int) -> str:
    """
    Wait until the result indicator has visible non-placeholder text, then
    return its text content (trimmed). Returns '' on timeout so the caller
    can decide to fail or warn.
    """
    LOG.debug("Waiting for result on %s", selector)
    deadline = time.monotonic() + timeout_s
    last = ""
    while time.monotonic() < deadline:
        txt = page.evaluate(
            "sel => { const el = document.querySelector(sel);"
            " if (!el) return '__MISSING__';"
            " const t = (el.innerText || el.textContent || '').trim();"
            " return t; }",
            arg=selector,
        )
        if txt and txt != "__MISSING__" and txt != "—":
            last = txt
            return txt
        page.wait_for_timeout(200)
    return last


# ──────────────────────────────────────────────────────────────────────────
# Per-tab driver
# ──────────────────────────────────────────────────────────────────────────


def xlsx_to_tsv(path: Path) -> str:
    """Read an XLSX file and return its first sheet as TSV text.

    Cells containing tabs, newlines, or non-breaking spaces are
    collapsed so they don't break the row/column structure that the
    WLP citizenship / PGR-training parsers expect.
    """
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        out_lines: list[str] = []
        for row in ws.iter_rows(values_only=True):
            if row is None:
                continue
            cells = []
            for c in row:
                if c is None:
                    cells.append("")
                else:
                    s = (
                        str(c)
                        .replace("\xa0", " ")
                        .replace("\t", " ")
                        .replace("\n", " ")
                        .replace("\r", " ")
                    )
                    s = " ".join(s.split())  # collapse runs of whitespace
                    cells.append(s)
            # Drop trailing empty cells so they don't look like extra columns
            while cells and cells[-1] == "":
                cells.pop()
            line = "\t".join(cells)
            if line.strip():
                out_lines.append(line)
        return "\n".join(out_lines)
    finally:
        wb.close()


def xlsx_sheets_to_tsvs(path: Path) -> list[str]:
    """Like xlsx_to_tsv() but returns one TSV per sheet (in workbook order)."""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        results: list[str] = []
        for sname in wb.sheetnames:
            ws = wb[sname]
            out_lines: list[str] = []
            for row in ws.iter_rows(values_only=True):
                if row is None:
                    continue
                cells = []
                for c in row:
                    if c is None:
                        cells.append("")
                    else:
                        s = (
                            str(c)
                            .replace("\xa0", " ")
                            .replace("\t", " ")
                            .replace("\n", " ")
                            .replace("\r", " ")
                        )
                        s = " ".join(s.split())
                        cells.append(s)
                while cells and cells[-1] == "":
                    cells.pop()
                line = "\t".join(cells)
                if line.strip():
                    out_lines.append(line)
            results.append("\n".join(out_lines))
        return results
    finally:
        wb.close()


def wait_for_citizenship_rows(page: Page, timeout_s: int) -> int:
    """Poll the citizenship tbody until it has at least one row, return count."""
    deadline = time.monotonic() + timeout_s
    while time.monotonic() < deadline:
        n = page.evaluate(
            "() => document.querySelectorAll('#citTbody tr').length"
        )
        if n and n > 0:
            return n
        page.wait_for_timeout(150)
    return 0


def wait_for_pgr_training_rows(page: Page, timeout_s: int) -> int:
    """Poll the PGR training tbody until it has at least one row."""
    deadline = time.monotonic() + timeout_s
    while time.monotonic() < deadline:
        n = page.evaluate(
            "() => document.querySelectorAll('#pgrtrTbody tr').length"
        )
        if n and n > 0:
            return n
        page.wait_for_timeout(150)
    return 0


def process_pgr_training_paste_tab(page: Page, tab: TabSpec, cfg: AppConfig) -> tuple[bool, str]:
    """Drive the PGR Training tab by converting each XLSX sheet into TSV
    and pasting it into the matching #pgrtr-tN textarea. A single click on
    #pgrtrAnalyseBtn then merges all sheets.
    """
    if not tab.files:
        return False, "No XLSX configured for PGR training paste"
    xlsx = tab.files[0]
    if not xlsx.exists():
        return False, f"Missing XLSX: {xlsx}"

    LOG.info("  (pgr training paste: %s)", xlsx.name)
    switch_main_tab(page, tab.main_tab)
    if tab.sub_tab:
        switch_sub_tab(page, tab.sub_tab)

    sheets = xlsx_sheets_to_tsvs(xlsx)
    textareas = ["#pgrtr-t1", "#pgrtr-t2", "#pgrtr-t3"]
    parts: list[str] = []
    for i, ta in enumerate(textareas):
        page.locator(ta).fill("")
        if i < len(sheets):
            tsv = sheets[i]
            row_count = tsv.count("\n") + 1 if tsv else 0
            if row_count > 0:
                LOG.info("  → pasting %d-row TSV into %s", row_count, ta)
                page.locator(ta).fill(tsv)
                parts.append(f"t{i+1}={row_count}r")
            else:
                LOG.info("  → sheet %d empty, leaving %s blank", i + 1, ta)

    if not parts:
        return False, "All three sheets were empty"

    LOG.info("  → clicking %s", tab.action)
    page.locator(tab.action).click()

    n_rows = wait_for_pgr_training_rows(page, cfg.result_timeout_s)
    if n_rows == 0:
        return False, "PGR training table did not populate"
    summary = f"{n_rows} rows ({', '.join(parts)})"

    if cfg.screenshot_dir:
        out = Path(cfg.screenshot_dir).expanduser().resolve()
        out.mkdir(parents=True, exist_ok=True)
        shot = out / f"{tab.key}.png"
        page.screenshot(path=str(shot), full_page=True)
        LOG.info("  📸 %s", shot)
    return True, summary


def process_citizenship_paste_tab(page: Page, tab: TabSpec, cfg: AppConfig) -> tuple[bool, str]:
    """Drive the School Citizenship tab via the textarea-paste path.

    The XLSX upload path overwrites on each upload, so to get all three
    categories in the combined citizenship table we paste TSV versions of
    each XLSX into the matching SharePoint-paste textarea and click
    Analyse once. The app's parseSharePointTable handles dedup across
    categories.
    """
    LOG.info("  (citizenship paste: %d categories)", len(tab.paste))
    switch_main_tab(page, tab.main_tab)
    if tab.sub_tab:
        switch_sub_tab(page, tab.sub_tab)

    textarea_map = {
        "teaching": "#cit-teaching",
        "research": "#cit-research",
        "school":   "#cit-school",
    }
    missing_inputs = [k for k in tab.paste if k not in textarea_map]
    if missing_inputs:
        return False, f"Unknown paste categories: {missing_inputs}"

    parts: list[str] = []
    for cat in ("teaching", "research", "school"):
        if cat not in tab.paste:
            page.locator(textarea_map[cat]).fill("")
            continue
        xlsx = tab.paste[cat]
        if not xlsx.exists():
            return False, f"Missing XLSX for {cat}: {xlsx}"
        tsv = xlsx_to_tsv(xlsx)
        row_count = tsv.count("\n") + 1 if tsv else 0
        LOG.info("  → pasting %d-row TSV into %s (%s)", row_count, textarea_map[cat], cat)
        page.locator(textarea_map[cat]).fill(tsv)
        parts.append(f"{cat}={row_count}r")

    LOG.info("  → clicking %s", tab.action)
    page.locator(tab.action).click()

    n_rows = wait_for_citizenship_rows(page, cfg.result_timeout_s)
    if n_rows == 0:
        return False, "Citizenship table did not populate"
    summary = f"{n_rows} combined rows ({', '.join(parts)})"

    if cfg.screenshot_dir:
        out = Path(cfg.screenshot_dir).expanduser().resolve()
        out.mkdir(parents=True, exist_ok=True)
        shot = out / f"{tab.key}.png"
        page.screenshot(path=str(shot), full_page=True)
        LOG.info("  📸 %s", shot)
    return True, summary


def process_combined_tag_setup_tab(
    page: Page, tab: TabSpec, cfg: AppConfig
) -> tuple[bool, str]:
    """Define a new tag rule on the Combined Totals tab and apply that tag to
    a specific academic.

    Requires that the combined table (#combTbody) is already populated — this
    step is meant to run *after* the `combined` tab (which clicks the Merge &
    Calculate button). The flow is:

      1. Make sure we are on the Combined Totals tab and the table has rows.
      2. Open the Tag Rules panel (it starts collapsed).
      3. Fill in the new-rule form (Tag name + Reduction %) and click Add.
      4. Find the row whose displayed name matches `apply_to_name` (normalised
         so "Dr Jonathan Burley" finds "Burley, Jonathan Dr" and the like).
      5. Click that row's '+ tag' button to open the tag popover.
      6. Type the tag name and click Add to apply it.

    Note: name matching does NOT collapse nicknames ("Jonathan" != "John"),
    unlike the WLP tool's own fuzzy matcher. Keep `apply_to_name` close to the
    spelling used in the data. If `apply_to_name` is empty, only the rule is
    defined — no row is matched and no tag is applied. `rule_reduction_pct`
    may be 0 (label-only tag) or positive (a workload allowance); negative
    values are rejected.
    """
    if not tab.rule_name or tab.rule_reduction_pct < 0:
        return False, "combined_tag_setup requires rule_name and rule_reduction_pct >= 0"

    LOG.info(
        "  (combined_tag_setup: rule=%r reduction=%d%% apply_to=%r)",
        tab.rule_name,
        tab.rule_reduction_pct,
        tab.apply_to_name,
    )

    # Make sure we're on the Combined tab.
    switch_main_tab(page, tab.main_tab)

    # 1. Wait for the combined table to actually have rows.
    deadline = time.monotonic() + cfg.result_timeout_s
    while time.monotonic() < deadline:
        n = page.evaluate("() => document.querySelectorAll('#combTbody tr').length")
        if n and n > 0:
            break
        page.wait_for_timeout(150)
    else:
        return False, "Combined table did not populate (run the 'combined' tab first?)"

    # 2. Open the rules panel body if it's collapsed (CSS display:none until .open).
    page.evaluate(
        "() => { const b = document.getElementById('rulesPanelBody');"
        " if (b && !b.classList.contains('open'))"
        " document.getElementById('rulesPanelHdr').click(); }"
    )
    page.wait_for_selector("#ruleNewTag", state="visible", timeout=5000)

    # 3. Fill the new-rule form and click Add.
    page.locator("#ruleNewTag").fill(tab.rule_name)
    page.locator("#ruleNewFte").fill(str(tab.rule_reduction_pct))
    page.locator("#ruleAddBtn").click()
    LOG.info("  → added tag rule %r with %d%% reduction", tab.rule_name, tab.rule_reduction_pct)

    # 3b. Optionally tick 'Teaching staff only' to narrow the visible table
    # before searching for the named academic. Idempotent.
    if tab.teaching_staff_only:
        page.evaluate(
            "() => { const cb = document.getElementById('combTlOnly');"
            " if (cb && !cb.checked) { cb.click(); return true; }"
            " return !!cb && cb.checked; }"
        )
        # Give the table a beat to re-render with the filter applied.
        page.wait_for_timeout(200)
        LOG.info("  → ticked 'Teaching staff only'")

    # 4-6. If `apply_to_name` is set, find the named academic, open the tag
    # popover, and apply the tag. Everything happens inside one
    # page.evaluate() rather than going through Playwright's locator.click()
    # because the popover is position:fixed and rendered near the anchor
    # button. The tool's openTagPopover() guesses the popover height as
    # ph=220px but the actual height is usually larger (Workload-allowance
    # row + Part-Time row + existing tags + suggestions), so the popover
    # overflows the viewport. Playwright's actionability checks then reject
    # the Add button as "outside of the viewport" and on later retries see it
    # as "not visible" (the document-level click handler closes the popover
    # on the spurious clicks those retries generate).
    #
    # Driving both the row button and the popover Add button via JS .click()
    # still goes through the real handlers (openTagPopover -> sets
    # tagPopoverCanonical; Add -> addTag + recomputeCombData + saveTagState +
    # combRender), so state and persistence are exactly what a human would get.
    # The match logic mirrors the WLP tool's `normaliseName()` (strip title
    # prefix, flip "Last, First" -> "First Last", lowercase, collapse
    # whitespace) so "Dr Jonathan Burley" finds "Burley, Jonathan Dr".
    applied_to: str | None = None
    if tab.apply_to_name:
        result = page.evaluate(
            """({applyToName, ruleName}) => {
                const norm = (s) => {
                  const TITLE_RE = /\\b(prof\\.?|professor|dr\\.?|mr\\.?|mrs\\.?|ms\\.?|mx\\.?|rev\\.?|sir)\\b\\s*/i;
                  let n = String(s || '').trim();
                  const cm = n.match(/^([^,]+),\\s*(.+)$/);
                  if (cm) n = cm[2] + ' ' + cm[1];
                  let prev;
                  do { prev = n; n = n.replace(TITLE_RE, ''); } while (n !== prev);
                  return n.toLowerCase().replace(/[-.']/g, ' ').replace(/\\s+/g, ' ').trim();
                };
                const rows = document.querySelectorAll('#combTbody tr');
                const needle = norm(applyToName);
                const needleTokens = needle.split(' ').filter(Boolean);
                let matched = null;
                for (const r of rows) {
                    const cn = r.querySelector('.cn');
                    if (!cn) continue;
                    const txt = norm(cn.textContent);
                    if (!txt) continue;
                    const txtTokens = txt.split(' ').filter(Boolean);
                    if (txt === needle) { matched = r; break; }
                    if (!matched) {
                        const isSub = txt.includes(needle) || needle.includes(txt);
                        const allTokensHit = needleTokens.length > 0 &&
                            needleTokens.every(nt => txtTokens.includes(nt));
                        if (isSub || allTokensHit) matched = r;
                    }
                }
                if (!matched) return {ok: false, reason: 'no row matched'};
                const btn = matched.querySelector('.comb-tag-add');
                if (!btn) return {ok: false, reason: 'no +tag button on row'};
                btn.click();
                const pop = document.getElementById('tagPopover');
                if (!pop || getComputedStyle(pop).display === 'none') {
                    return {ok: false, reason: 'popover did not open'};
                }
                const input = document.getElementById('tagPopoverInput');
                const addBtn = document.getElementById('tagPopoverAdd');
                if (!input || !addBtn) return {ok: false, reason: 'popover inputs missing'};
                input.value = ruleName;
                addBtn.click();
                return {ok: true, display: matched.querySelector('.cn').textContent.trim()};
            }""",
            arg={"applyToName": tab.apply_to_name, "ruleName": tab.rule_name},
        )

        if not result.get("ok"):
            sample = page.evaluate(
                "() => [...document.querySelectorAll('#combTbody .cn')].slice(0,15).map(t => t.textContent.trim())"
            )
            return False, (
                f"Tag application failed ({result.get('reason', 'unknown')}). "
                f"Sample names: {sample}"
            )

        applied_to = result["display"]
        LOG.info("  → applied tag %r to %s", tab.rule_name, applied_to)

        # Tidy up: close the popover (harmless if already closed, but cleaner).
        page.evaluate(
            "() => { const p = document.getElementById('tagPopover');"
            " if (p) p.style.display = 'none'; }"
        )
    else:
        LOG.info(
            "  → rule %r defined (no apply_to_name, no row tagged)",
            tab.rule_name,
        )

    if applied_to:
        summary = (
            f"rule {tab.rule_name!r} ({tab.rule_reduction_pct}% reduction) + "
            f"applied to {applied_to!r}"
        )
    else:
        summary = (
            f"rule {tab.rule_name!r} ({tab.rule_reduction_pct}% reduction) defined"
        )
    if cfg.screenshot_dir:
        out = Path(cfg.screenshot_dir).expanduser().resolve()
        out.mkdir(parents=True, exist_ok=True)
        shot = out / f"{tab.key}.png"
        page.screenshot(path=str(shot), full_page=True)
        LOG.info("  📸 %s", shot)
    return True, summary


def process_combined_tag_assignments_tab(
    page: Page, tab: TabSpec, cfg: AppConfig
) -> tuple[bool, str]:
    """Apply tags to many academics at once from an XLSX file.

    The XLSX must have:
      * one row per academic (header on row 1),
      * the academic's name in column `assignments_name_column` (1-indexed,
        default 1),
      * one column per tag whose header is the tag name — listed in
        `assignments_tag_columns` (tag-name -> column-index).

    If the cell value matches `assignments_apply_marker` (default "Y",
    case-insensitive after trimming), that tag is applied to the academic.

    For each (name, [tags]) row the script:
      1. finds the academic's row in #combTbody (normalised match, same
         logic as `combined_tag_setup_tab`),
      2. clicks that row's '+ tag' button — this opens the tag popover and
         sets tagPopoverCanonical,
      3. types and Adds each tag inside the popover,
      4. closes the popover before moving on.

    All of one person's tags go in a single popover session, but each
    person still needs its own row-find + popover-open, so this is N
    page.evaluate() calls for N people. The matcher uses the same
    normaliser as `combined_tag_setup_tab` so "Dr Jonathan Burley" still
    finds "Burley, Jonathan Dr".

    Tag rules must already exist (defined by a preceding
    `combined_tag_setup` entry) for the tag to take effect on FTE. This
    step only stores the tag on the academic via `addTag()`; without a
    matching rule the tag is a label-only.
    """
    if not tab.assignments_xlsx:
        return False, "combined_tag_assignments requires assignments_xlsx"
    if not tab.assignments_tag_columns:
        return False, "combined_tag_assignments requires assignments_tag_columns"

    xlsx_path = Path(tab.assignments_xlsx).expanduser()
    if not xlsx_path.exists():
        return False, f"XLSX not found: {xlsx_path}"

    # 1. Read the XLSX and build the assignments list.
    try:
        wb = load_workbook(str(xlsx_path), read_only=True, data_only=True)
    except Exception as e:
        return False, f"Could not open {xlsx_path}: {e}"

    try:
        ws = (
            wb[tab.assignments_sheet]
            if tab.assignments_sheet
            else wb[wb.sheetnames[0]]
        )
        marker = tab.assignments_apply_marker.strip().upper()
        assignments: list[tuple[str, list[str]]] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or tab.assignments_name_column > len(row):
                continue
            raw_name = row[tab.assignments_name_column - 1]
            if not raw_name or not isinstance(raw_name, str):
                continue
            tags: list[str] = []
            for tag_name, col_idx in tab.assignments_tag_columns.items():
                if col_idx > len(row):
                    continue
                cell = row[col_idx - 1]
                if cell is not None and str(cell).strip().upper() == marker:
                    tags.append(tag_name)
            if tags:
                assignments.append((raw_name.strip(), tags))
    finally:
        wb.close()

    if not assignments:
        return True, "No assignments found in XLSX (nothing to apply)"

    LOG.info(
        "  (combined_tag_assignments: %d people, %d tags total)",
        len(assignments),
        sum(len(t) for _, t in assignments),
    )

    # 2. Make sure we're on the Combined tab and the table has rows.
    switch_main_tab(page, tab.main_tab)
    deadline = time.monotonic() + cfg.result_timeout_s
    while time.monotonic() < deadline:
        n = page.evaluate("() => document.querySelectorAll('#combTbody tr').length")
        if n and n > 0:
            break
        page.wait_for_timeout(150)
    else:
        return False, "Combined table did not populate (run the 'combined' tab first?)"

    # 3. For each (name, tags): find the row, open the popover, Add each tag.
    applied_tags = 0
    applied_people = 0
    failures: list[tuple[str, str]] = []
    for name, tags in assignments:
        result = page.evaluate(
            """({applyToName, tags}) => {
                const norm = (s) => {
                  const TITLE_RE = /\\b(prof\\.?|professor|dr\\.?|mr\\.?|mrs\\.?|ms\\.?|mx\\.?|rev\\.?|sir)\\b\\s*/i;
                  let n = String(s || '').trim();
                  const cm = n.match(/^([^,]+),\\s*(.+)$/);
                  if (cm) n = cm[2] + ' ' + cm[1];
                  let prev;
                  do { prev = n; n = n.replace(TITLE_RE, ''); } while (n !== prev);
                  return n.toLowerCase().replace(/[-.']/g, ' ').replace(/\\s+/g, ' ').trim();
                };
                const rows = document.querySelectorAll('#combTbody tr');
                const needle = norm(applyToName);
                const needleTokens = needle.split(' ').filter(Boolean);
                let matched = null;
                for (const r of rows) {
                    const cn = r.querySelector('.cn');
                    if (!cn) continue;
                    const txt = norm(cn.textContent);
                    if (!txt) continue;
                    const txtTokens = txt.split(' ').filter(Boolean);
                    if (txt === needle) { matched = r; break; }
                    if (!matched) {
                        const isSub = txt.includes(needle) || needle.includes(txt);
                        const allTokensHit = needleTokens.length > 0 &&
                            needleTokens.every(nt => txtTokens.includes(nt));
                        if (isSub || allTokensHit) matched = r;
                    }
                }
                if (!matched) return {ok: false, reason: 'no row matched'};
                const btn = matched.querySelector('.comb-tag-add');
                if (!btn) return {ok: false, reason: 'no +tag button'};
                btn.click();
                const pop = document.getElementById('tagPopover');
                if (!pop || getComputedStyle(pop).display === 'none') {
                    return {ok: false, reason: 'popover did not open'};
                }
                const input = document.getElementById('tagPopoverInput');
                const addBtn = document.getElementById('tagPopoverAdd');
                if (!input || !addBtn) return {ok: false, reason: 'popover inputs missing'};
                const applied = [];
                for (const tag of tags) {
                    input.value = tag;
                    addBtn.click();
                    applied.push(tag);
                }
                pop.style.display = 'none';
                return {ok: true, display: matched.querySelector('.cn').textContent.trim(), applied};
            }""",
            arg={"applyToName": name, "tags": tags},
        )
        if result.get("ok"):
            applied_tags += len(result.get("applied", []))
            applied_people += 1
        else:
            failures.append((name, result.get("reason", "unknown")))

    LOG.info("  → applied %d tags to %d people", applied_tags, applied_people)
    for fname, freason in failures[:5]:
        LOG.warning("  ⚠ %s: %s", fname, freason)
    if len(failures) > 5:
        LOG.warning("  ⚠ (%d more unmatched)", len(failures) - 5)

    if applied_tags == 0 and failures:
        return False, (
            f"No tags applied; first failure: {failures[0][0]!r} ({failures[0][1]}). "
            f"{len(failures)} people unmatched in total."
        )

    summary_parts = [f"applied {applied_tags} tags to {applied_people} people"]
    if failures:
        summary_parts.append(
            f"{len(failures)} unmatched (first: {failures[0][0]!r})"
        )

    if cfg.screenshot_dir:
        out = Path(cfg.screenshot_dir).expanduser().resolve()
        out.mkdir(parents=True, exist_ok=True)
        shot = out / f"{tab.key}.png"
        page.screenshot(path=str(shot), full_page=True)
        LOG.info("  📸 %s", shot)

    return True, "; ".join(summary_parts)


def process_tab(page: Page, tab: TabSpec, cfg: AppConfig) -> tuple[bool, str]:
    """Run one tab end-to-end. Returns (success, summary_message)."""
    LOG.info("─" * 64)
    LOG.info("Tab: %s", tab.key)
    if tab.notes.strip():
        LOG.info("  (%s)", " ".join(tab.notes.split()))

    # Branch on kind.
    if tab.kind == "citizenship_paste":
        ok, summary = process_citizenship_paste_tab(page, tab, cfg)
    elif tab.kind == "pgr_training_paste":
        ok, summary = process_pgr_training_paste_tab(page, tab, cfg)
    elif tab.kind == "combined_tag_setup":
        ok, summary = process_combined_tag_setup_tab(page, tab, cfg)
    elif tab.kind == "combined_tag_assignments":
        ok, summary = process_combined_tag_assignments_tab(page, tab, cfg)
    else:
        ok, summary = process_upload_tab(page, tab, cfg)

    if ok:
        LOG.info("  ✓ %s", summary[:160])
    return ok, summary


def process_upload_tab(page: Page, tab: TabSpec, cfg: AppConfig) -> tuple[bool, str]:
    """Default upload path: pick files, click action, wait for wait_for.

    If `tab.file_input` is empty, the upload step is skipped (useful for
    tabs like Combined that just trigger an action on already-loaded data).
    """
    # 1. Switch to the tab.
    switch_main_tab(page, tab.main_tab)
    if tab.sub_tab:
        switch_sub_tab(page, tab.sub_tab)

    # 2. Validate file paths.
    missing = [f for f in tab.files if not f.exists()]
    if missing:
        msg = f"Missing data files: {[str(m) for m in missing]}"
        LOG.error("  ✗ %s", msg)
        return False, msg

    # 3. Upload via Playwright's native mechanism. This fires the page's
    #    'change' event exactly like a human picking a file. Skipped if
    #    no file_input is configured.
    if tab.file_input and tab.files:
        file_strs = [str(f) for f in tab.files]
        LOG.info("  → uploading %d file(s) into %s", len(file_strs), tab.file_input)
        page.locator(tab.file_input).set_input_files(file_strs)
    elif not tab.file_input:
        LOG.info("  → no file input (action-only tab)")
    else:
        LOG.info("  → no files configured for this tab")

    # 3b. Preflight: tick boxes / click elements BEFORE the action button.
    if tab.preflight:
        run_steps(page, tab.preflight, "preflight")

    # 4. Click the action button if there is one.
    if tab.action:
        try:
            wait_for_action_enabled(page, tab.action, timeout_s=10)
        except PWTimeoutError as e:
            msg = f"Action button {tab.action} never became enabled"
            LOG.error("  ✗ %s", msg)
            return False, msg
        LOG.info("  → clicking %s", tab.action)
        # Click via JS .click() to avoid Playwright's actionability/idle
        # waits that can hang on heavy pages (e.g. the 2 kLOC doMerge).
        try:
            clicked = page.evaluate(
                "sel => { const el=document.querySelector(sel);"
                " if(!el)return false; el.click(); return true; }",
                arg=tab.action,
            )
            if not clicked:
                msg = f"Could not find {tab.action} to click"
                LOG.error("  ✗ %s", msg)
                return False, msg
        except Exception as e:
            msg = f"Click on {tab.action} failed: {e}"
            LOG.error("  ✗ %s", msg)
            return False, msg

        # 4b. Postflight: tick boxes / click elements AFTER the action
        # button (e.g. the Teaching tab's #tlRecalcBtn to apply prep hours
        # once the analyser page is rendered).
        if tab.postflight:
            # Give the page a moment to re-render after the action click
            # before running postflight steps.
            page.wait_for_timeout(300)
            run_steps(page, tab.postflight, "postflight")
    else:
        LOG.info("  → no action button (results render automatically)")

    # 5. Wait for the result indicator to populate.
    summary = wait_for_result(page, tab.wait_for, cfg.result_timeout_s)
    if not summary:
        msg = f"Result indicator {tab.wait_for} did not populate within {cfg.result_timeout_s}s"
        LOG.error("  ✗ %s", msg)
        return False, msg
    one_line = " ".join(summary.split())[:160]

    # 5b. For the Combined tab, also print the first 20 displayed academic
    # names so the user can verify the name reformatting visually.
    if tab.key == "combined":
        debug = page.evaluate("""() => {
            const rows = document.querySelectorAll('#combTbody tr');
            const cn = document.querySelectorAll('#combTbody td.cn');
            const allTds = document.querySelectorAll('#combTbody td');
            return {
                row_count: rows.length,
                cn_count: cn.length,
                td_count: allTds.length,
                first_row_html: rows[0] ? rows[0].outerHTML.slice(0, 800) : 'none',
                names: [...cn].slice(0,20).map(td => td.innerText.trim().split('\\n')[0]),
            };
        }""")
        LOG.info("  Combined debug: rows=%d td.cn=%d td.total=%d",
                 debug["row_count"], debug["cn_count"], debug["td_count"])
        if debug["first_row_html"] != "none":
            LOG.info("  First row HTML: %s", debug["first_row_html"][:600])
        if debug["names"]:
            LOG.info("  Combined view — first 20 names:")
            for n in debug["names"]:
                LOG.info("    %s", n)

    # 6. Optional screenshot.
    if cfg.screenshot_dir:
        out = Path(cfg.screenshot_dir).expanduser().resolve()
        out.mkdir(parents=True, exist_ok=True)
        shot = out / f"{tab.key}.png"
        page.screenshot(path=str(shot), full_page=True)
        LOG.info("  📸 %s", shot)

    return True, one_line


# ──────────────────────────────────────────────────────────────────────────
# Entry point
# ──────────────────────────────────────────────────────────────────────────


def setup_logging() -> None:
    fmt = "%(asctime)s %(levelname)-5s %(message)s"
    logging.basicConfig(level=logging.INFO, format=fmt, datefmt="%H:%M:%S")


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    p.add_argument(
        "--config",
        default=str(Path(__file__).parent / "config" / "tabs.yaml"),
        help="Path to YAML config (default: config/tabs.yaml next to this script)",
    )
    p.add_argument(
        "--dry-run",
        action="store_true",
        help="Print the planned sequence without launching a browser",
    )
    return p.parse_args()


def main() -> int:
    setup_logging()
    args = parse_args()

    cfg_path = Path(args.config)
    if not cfg_path.exists():
        LOG.error("Config not found: %s", cfg_path)
        return 2
    cfg = AppConfig.load(cfg_path)
    enabled = [t for t in cfg.tabs if t.enabled]
    LOG.info("Loaded %d tab(s) from %s (%d enabled)", len(cfg.tabs), cfg_path, len(enabled))
    LOG.info("Target URL: %s", cfg.url)
    LOG.info("Data dir:   %s", cfg.data_dir)
    for t in enabled:
        LOG.info("  • %s — %d file(s)", t.key, len(t.files))

    if args.dry_run:
        LOG.info("--dry-run: not launching browser")
        return 0

    results: list[tuple[str, bool, str]] = []
    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=cfg.headless,
            slow_mo=cfg.slow_mo_ms,
            args=["--no-sandbox", "--disable-dev-shm-usage"],
        )
        context = browser.new_context(viewport={"width": 1400, "height": 900})
        page = context.new_page()
        page.set_default_timeout(15000)

        LOG.info("Opening %s", cfg.url)
        page.goto(cfg.url, wait_until="domcontentloaded")
        # The app initialises several things on DOMContentLoaded; give it a beat.
        page.wait_for_selector(".main-tab-btn", timeout=20000)
        dismiss_session_banner(page)

        for tab in enabled:
            ok, summary = process_tab(page, tab, cfg)
            results.append((tab.key, ok, summary))

        # Summary banner in the terminal.
        LOG.info("═" * 64)
        LOG.info("Summary")
        for key, ok, summary in results:
            mark = "✓" if ok else "✗"
            LOG.info("  %s %-26s %s", mark, key, summary)
        ok_count = sum(1 for _, ok, _ in results if ok)
        LOG.info("  %d/%d succeeded", ok_count, len(results))

        # Hold the browser open so the user can inspect.
        LOG.info("Browser is still open. Press Enter to close and exit.")
        try:
            input()
        except EOFError:
            # Non-interactive (e.g. CI) — just sleep briefly so the user
            # has a chance to glance, then close.
            page.wait_for_timeout(2000)

        context.close()
        browser.close()

    return 0 if ok_count == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())
